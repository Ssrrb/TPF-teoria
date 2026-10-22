#!/usr/bin/env python3
"""modelo.yaml + modelo/out/*.png -> TPF-Obra-Studio.docx
Reproduce exactamente la estructura de la plantilla de la cátedra."""
import pathlib, struct, yaml
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

RAIZ = pathlib.Path(__file__).resolve().parent.parent
PROY = RAIZ.parent
from modelo import cargar
M = cargar()
OUT = RAIZ / "out"

meta     = M["meta"]
actores  = {a["id"]: a for a in M["actores"]}
casos    = {c["id"]: c for c in M["casos_uso"]}
reqs     = M["requisitos"]
detallados = [c for c in M["casos_uso"] if c.get("detallado")]

SERIF, SANS = "Liberation Serif", "Liberation Sans"
ANCHO_V = Cm(16.0)          # ancho útil en página vertical
ANCHO_H_MAX = 25.0          # cm de ancho útil en página horizontal
ALTO_H_MAX = 16.3           # cm de alto realmente utilizable (medido sobre la salida)
ALTO_TITULO, ALTO_PIE = 1.2, 0.6

def png_dim(p):
    d = open(p, "rb").read(33)
    return struct.unpack(">II", d[16:24])

# ------------------------------------------------------------------ documento
doc = Document()
st = doc.styles["Normal"]
st.font.name = SERIF; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), SERIF)
st.paragraph_format.space_after = Pt(6)
for n, sz in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
    s = doc.styles[n]
    s.font.name = SANS; s.font.size = Pt(sz); s.font.bold = True
    s.font.color.rgb = None
    s.element.rPr.rFonts.set(qn("w:eastAsia"), SANS)

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for m_ in ("top_margin", "bottom_margin"): setattr(sec, m_, Cm(2.2))
for m_ in ("left_margin", "right_margin"): setattr(sec, m_, Cm(2.5))
sec.different_first_page_header_footer = True

CAB = [meta["institucion"], meta["facultad"], meta["departamento"],
       meta["catedra"], f'Prof.: {meta["docente"]}']

def poner_cabecera(seccion, primera_distinta=False):
    seccion.different_first_page_header_footer = primera_distinta
    enc = seccion.header
    enc.is_linked_to_previous = False
    for p in list(enc.paragraphs)[1:]:
        p._element.getparent().remove(p._element)
    p = enc.paragraphs[0]; p.text = ""
    for i, linea in enumerate(CAB):
        r = p.add_run(("\n" if i else "") + linea)
        r.font.name = SANS; r.font.size = Pt(8)
    p.paragraph_format.space_after = Pt(0)

def numero_de_pagina(seccion):
    seccion.footer.is_linked_to_previous = False
    p = seccion.footer.paragraphs[0]
    p.text = ""; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(); r.font.name = SERIF; r.font.size = Pt(9)
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = "PAGE"
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    r._r.append(f1); r._r.append(it); r._r.append(f2)

ESTADO = {"orient": "P"}

def _nueva_seccion(horizontal):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    if horizontal:
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Cm(29.7), Cm(21)
        s.top_margin, s.bottom_margin = Cm(2.4), Cm(1.2)   # el encabezado ocupa 5 líneas
        for m_ in ("left_margin", "right_margin"): setattr(s, m_, Cm(1.8))
    else:
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        for m_ in ("top_margin", "bottom_margin"): setattr(s, m_, Cm(2.2))
        for m_ in ("left_margin", "right_margin"): setattr(s, m_, Cm(2.5))
    poner_cabecera(s); numero_de_pagina(s)
    return s

def ir_horizontal():
    if ESTADO["orient"] != "L":
        _nueva_seccion(True); ESTADO["orient"] = "L"

def ir_vertical():
    if ESTADO["orient"] != "P":
        _nueva_seccion(False); ESTADO["orient"] = "P"

def par(texto="", estilo=None, align=None, tam=None, bold=False, space=6, bruto=False):
    if not bruto: ir_vertical()
    p = doc.add_paragraph(style=estilo)
    if texto:
        r = p.add_run(texto); r.bold = bold
        if tam: r.font.size = Pt(tam)
    if align is not None: p.alignment = align
    p.paragraph_format.space_after = Pt(space)
    return p

def h(nivel, texto, bruto=False):
    if not bruto: ir_vertical()
    return doc.add_heading(texto, level=nivel)

def imagen(nombre, pie, titulo=None, intro=None):
    """Coloca el diagrama; los muy apaisados van en una página horizontal.
    El título y la introducción acompañan al diagrama para no dejar páginas casi vacías."""
    ruta = OUT / nombre
    pw, ph = png_dim(ruta)
    ratio = pw / ph
    # a 16 cm un diagrama grande queda ilegible: va a página horizontal
    apaisado = ratio > 1.62 or max(pw, ph) > 2200
    if apaisado:
        ir_horizontal()
        if titulo: h(titulo[0], titulo[1], bruto=True)
        if intro: par(intro, tam=10, bruto=True)
        # el diagrama ocupa lo que queda tras el título y el pie de figura
        alto_disp = ALTO_H_MAX - (ALTO_TITULO if titulo else 0) - ALTO_PIE
        if intro: alto_disp -= 1.2
        if ratio >= ANCHO_H_MAX / alto_disp:
            ancho, alto = Cm(ANCHO_H_MAX), None
        else:
            ancho, alto = None, Cm(alto_disp)
    else:
        ir_vertical()
        if titulo: h(titulo[0], titulo[1])
        if intro: par(intro, tam=10)
        ancho, alto = ANCHO_V, None
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if ancho is not None:
        p.add_run().add_picture(str(ruta), width=ancho)
    else:
        p.add_run().add_picture(str(ruta), height=alto)
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(pie); r.font.size = Pt(9); r.italic = True

def tabla(filas, anchos=None, encabezado=False, tam=9.5, bruto=False):
    if not bruto: ir_vertical()
    t = doc.add_table(rows=0, cols=len(filas[0])); t.style = "Table Grid"
    for i, fila in enumerate(filas):
        celdas = t.add_row().cells
        for j, v in enumerate(fila):
            celdas[j].text = ""
            p = celdas[j].paragraphs[0]
            r = p.add_run(str(v))
            r.font.size = Pt(tam); r.font.name = SERIF
            if encabezado and i == 0: r.bold = True
            p.paragraph_format.space_after = Pt(2)
            if anchos: celdas[j].width = anchos[j]
    if anchos:
        for fila in t.rows:
            for j, c in enumerate(fila.cells): c.width = anchos[j]
    return t

# --------------------------------------------------------------------- portada
par(meta["facultad"], align=WD_ALIGN_PARAGRAPH.CENTER, tam=14, bold=True, space=2)
par(meta["carrera"], align=WD_ALIGN_PARAGRAPH.CENTER, tam=14, bold=True, space=2)
par(meta["materia"], align=WD_ALIGN_PARAGRAPH.CENTER, tam=14, bold=True, space=2)
par(meta["trabajo"], align=WD_ALIGN_PARAGRAPH.CENTER, tam=14, bold=True, space=60)
par("Autor/es", align=WD_ALIGN_PARAGRAPH.CENTER, tam=12, bold=True, space=4)
for a in meta["autores"]:
    par(a["nombre"], align=WD_ALIGN_PARAGRAPH.CENTER, tam=12, space=2)
par("", space=48)
par("Profesor", align=WD_ALIGN_PARAGRAPH.CENTER, tam=12, bold=True, space=4)
par(meta["docente"], align=WD_ALIGN_PARAGRAPH.CENTER, tam=12, space=60)
par(meta["ciudad"], align=WD_ALIGN_PARAGRAPH.CENTER, tam=12, space=2)
par(meta["fecha"], align=WD_ALIGN_PARAGRAPH.CENTER, tam=12)
doc.add_page_break()
poner_cabecera(sec, primera_distinta=True); numero_de_pagina(sec)

# -------------------------------------------------------------------- contenido
par("Contenido", align=WD_ALIGN_PARAGRAPH.CENTER, tam=14, bold=True, space=12)
indice = [(0, "Modelo de Análisis"), (1, "Requerimientos"),
          (2, "Requerimientos Funcionales"), (2, "Requerimientos No Funcionales"),
          (1, "Identificación de Actores"), (1, "Diagrama de Casos de Uso")]
for sid, s in sorted(M["subsistemas"].items()):
    indice.append((2, f'Sistema {meta["sistema"]} — Subsistema {s["nombre"]}'))
for c in detallados:
    indice.append((1, f'Caso de Uso {c["id"]} {c["nombre"]}'))
    indice += [(2, "Especificación Caso de Uso"), (2, "Diagrama de Clases del Caso de Uso"),
               (2, "Diagrama de Interacción del Caso de Uso")]
indice += [(1, "Matriz de Requerimientos vs Casos de Uso"), (0, "Modelo de Diseño"),
           (1, "Diagrama de Clases Persistentes"), (1, "Diagrama de Entidad Relación"),
           (0, "Referencias")]
for nivel, texto in indice:
    p = par(texto, tam=11, bold=(nivel == 0), space=2)
    p.paragraph_format.left_indent = Cm(0.8 * nivel)
doc.add_page_break()

# ============================================================ MODELO DE ANÁLISIS
h(1, "Modelo de Análisis")
par(f'Este modelo describe el análisis de {meta["sistema"]}, sistema de gestión de obra que separa '
    f'dos motores: el ECMS mide la obra y el ERP registra lo que el ECMS midió. El alcance desarrollado '
    f'comprende presupuesto, análisis de costo unitario, certificación y subcontratos, con facturas y '
    f'pagos tratados como interacción con el ERP. La arquitectura de referencia es Clean Architecture; '
    f'sus anillos se proyectan sobre la notación de análisis mediante los estereotipos «Boundary» '
    f'(adaptadores de interfaz), «Control» (casos de uso) y «Entity» (dominio), y los puertos aparecen '
    f'como «Interface» realizados por la infraestructura.')

h(2, "Requerimientos")

h(3, "Requerimientos Funcionales")
filas = [["ID", "Requerimiento"]]
filas += [[r["id"], " ".join(r["enunciado"].split())] for r in reqs if r["tipo"] == "RF"]
tabla(filas, anchos=[Cm(1.8), Cm(14.2)], encabezado=True)

h(3, "Requerimientos No Funcionales")
filas = [["ID", "Requerimiento"]]
filas += [[r["id"], " ".join(r["enunciado"].split())] for r in reqs if r["tipo"] in ("RNF", "PR")]
tabla(filas, anchos=[Cm(1.8), Cm(14.2)], encabezado=True)
par("Los identificadores PR corresponden a restricciones de plataforma y alcance.", tam=9)

h(2, "Identificación de Actores")
filas = [["Actor", "Descripción"]]
filas += [[a["nombre"], " ".join(a["descripcion"].split())] for a in M["actores"]]
tabla(filas, anchos=[Cm(4.0), Cm(12.0)], encabezado=True)

h(2, "Diagrama de Casos de Uso")
for i, (sid, s) in enumerate(sorted(M["subsistemas"].items()), start=1):
    imagen(f'0{i}-casos-de-uso-{sid}.png',
           f'Figura {i}. Casos de uso del subsistema {s["nombre"]}.',
           titulo=(3, f'Sistema {meta["sistema"]} — Subsistema {s["nombre"]}'),
           intro=s["descripcion"])

# ------------------------------------------------------- especificación de un CU
def spec_tabla(c):
    A = [Cm(3.4), Cm(6.0), Cm(2.2), Cm(4.4)]
    t = doc.add_table(rows=0, cols=4); t.style = "Table Grid"

    def fila(valores, combinar=None, negrita_0=True, centrar=False, tam=9.5):
        celdas = t.add_row().cells
        if combinar:
            ini, fin = combinar
            celdas[ini].merge(celdas[fin])
            celdas = t.rows[-1].cells
        for j, v in enumerate(valores):
            if j >= len(celdas): break
            celdas[j].text = ""
            p = celdas[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            if centrar: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for k, linea in enumerate(str(v).split("\n")):
                r = p.add_run(("\n" if k else "") + linea)
                r.font.size = Pt(tam); r.font.name = SERIF
                r.bold = (j == 0 and negrita_0) or centrar
        for j, cel in enumerate(t.rows[-1].cells):
            if j < len(A): cel.width = A[j]

    fila(["Nombre Caso de Uso", f'{c["id"]} {c["nombre"]}', "Versión", meta["version"]])
    fila(["Creado por", c["creado_por"], "Fecha", c["fecha_creacion"]])
    fila(["Modificado por", c["modificado_por"], "Fecha", c["fecha_modificacion"]])
    fila(["Descripción", " ".join(c["descripcion"].split())], combinar=(1, 3))
    reqs_cu = [r["id"] for r in reqs if c["id"] in (r.get("casos_uso") or [])]
    fila(["Requerimiento(s) a que Responde", ", ".join(reqs_cu)], combinar=(1, 3))
    fila(["Actores", ", ".join(actores[a]["nombre"] for a in c["actores"])], combinar=(1, 3))
    fila(["Pre Condiciones", "\n".join(f"· {x}" for x in c["precondiciones"])], combinar=(1, 3))
    fila(["Pos Condiciones", "\n".join(f"· {x}" for x in c["poscondiciones"])], combinar=(1, 3))
    fila(["Flujo Normal"], combinar=(0, 3), centrar=True)
    fila(["Actor", "Acción", "Respuesta del Sistema"], combinar=(2, 3), centrar=True)
    for paso in c["flujo"]:
        fila([actores[paso["actor"]]["nombre"] if paso["actor"] in actores else paso["actor"],
              paso["accion"], paso["respuesta"]], combinar=(2, 3), negrita_0=False)
    fila(["Excepciones",
          "\n".join(f'{e["id"]} ({e["requisito"]}). {e["texto"]}' for e in c["excepciones"])],
         combinar=(1, 3))
    return t

fig = len(M["subsistemas"])
for c in detallados:
    h(2, f'Caso de Uso {c["id"]} {c["nombre"]}')
    h(3, "Especificación Caso de Uso")
    spec_tabla(c)
    fig += 1
    n = {"CU-02": "03", "CU-05": "04", "CU-09": "05"}[c["id"]]
    imagen(f'{n}-{c["id"]}-clases.png', f'Figura {fig}. Clases del caso de uso {c["id"]}.',
           titulo=(3, "Diagrama de Clases del Caso de Uso"))
    fig += 1
    n = {"CU-02": "06", "CU-05": "07", "CU-09": "08"}[c["id"]]
    imagen(f'{n}-{c["id"]}-secuencia.png', f'Figura {fig}. Interacción del caso de uso {c["id"]}.',
           titulo=(3, "Diagrama de Interacción del Caso de Uso"))

# ------------------------------------------------ matriz requerimientos vs CU
ir_horizontal()
h(2, "Matriz de Requerimientos vs Casos de Uso", bruto=True)
par(" · ".join(f'{c["id"]} {c["nombre"]}' for c in M["casos_uso"]), tam=8.5, bruto=True)

ids_cu = [c["id"] for c in M["casos_uso"]]
filas = [["Requerimiento"] + [c.split("-")[1] for c in ids_cu]]
for r in reqs:
    marca = set(r.get("casos_uso") or [])
    filas.append([r["id"]] + ["●" if c in marca else "" for c in ids_cu])
anchos = [Cm(2.6)] + [Cm(1.33)] * len(ids_cu)
t = doc.add_table(rows=0, cols=len(filas[0])); t.style = "Table Grid"
for i, fila in enumerate(filas):
    celdas = t.add_row().cells
    for j, v in enumerate(fila):
        celdas[j].text = ""; celdas[j].width = anchos[j]
        p = celdas[j].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        if j: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(v)); run.font.size = Pt(8); run.font.name = SERIF
        run.bold = (i == 0)
for fila in t.rows:
    for j, c_ in enumerate(fila.cells): c_.width = anchos[j]
trPr = t.rows[0]._tr.get_or_add_trPr()
th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
par(bruto=True, texto=f'Los {len(reqs)} identificadores del catálogo quedan cubiertos: '
    f'{sum(1 for r in reqs if r["tipo"] == "RF")} funcionales, '
    f'{sum(1 for r in reqs if r["tipo"] == "RNF")} no funcionales y '
    f'{sum(1 for r in reqs if r["tipo"] == "PR")} restricciones de plataforma. '
    f'Las columnas corresponden a los casos de uso CU-01 a CU-{len(ids_cu):02d}.', tam=9)

# ============================================================= MODELO DE DISEÑO
h(1, "Modelo de Diseño")

h(2, "Diagrama de Clases Persistentes")
derivados = [f'{e["nombre"]}.{a["n"]}' for e in M["entidades"]
             for a in e["atributos"] if a.get("derivado")]
par(f'Las clases persistentes son {sum(1 for e in M["entidades"] if e.get("tabla"))}, agrupadas en '
    f'paquetes y repartidas en dos figuras complementarias para que resulten legibles. Los '
    f'{len(derivados)} atributos marcados con una barra inclinada son derivados: se calculan a '
    f'partir de los registros aprobados y no se almacenan, de modo que no puedan discrepar de su '
    f'origen. Los puntos suspensivos indican atributos elididos; el detalle completo de columnas y '
    f'tipos figura en el diagrama de entidad relación. Valorizacion es un objeto de valor y no '
    f'persiste por sí mismo.')
for num, sid, nombre in [("09", "A", "obra y presupuesto"), ("10", "B", "medición y certificación")]:
    fig += 1
    imagen(f'{num}-clases-persistentes-{sid}.png',
           f'Figura {fig}. Clases persistentes ({sid}): {nombre}.')

h(2, "Diagrama de Entidad Relación")
tablas_n = sum(1 for e in M["entidades"] if e.get("tabla"))
par(f'El modelo relacional comprende {tablas_n} tablas y {len(M["relaciones"])} relaciones, en '
    f'notación de pata de gallo. El símbolo ● identifica una columna obligatoria; «PK» y «FK» '
    f'señalan la clave primaria y las claves foráneas. Los atributos derivados del diagrama anterior '
    f'no aparecen como columnas.')
fig += 1
imagen("11-entidad-relacion.png", f"Figura {fig}. Diagrama de entidad relación.")

# ================================================================= REFERENCIAS
h(1, "Referencias")
for ref in M["referencias"]:
    texto = ref.replace("*", "")
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.first_line_indent = Cm(-1.0)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(texto); r.font.size = Pt(10.5); r.font.name = SERIF

destino = PROY / "TPF-Obra-Studio.docx"
doc.save(str(destino))
print(f"Documento generado: {destino}")
print(f"  {len(reqs)} requisitos · {len(M['casos_uso'])} casos de uso · "
      f"{len(detallados)} especificados · {fig} figuras")
